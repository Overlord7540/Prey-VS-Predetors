"""Build Predator_Prey_Lab_Presentation.pptx from slide content + screenshots."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUT = ROOT / "Predator_Prey_Lab_Presentation.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x0F, 0x14, 0x19)
PANEL = RGBColor(0x1A, 0x23, 0x32)
TEXT = RGBColor(0xE8, 0xEE, 0xF4)
MUTED = RGBColor(0x8B, 0x9A, 0xAB)
ACCENT = RGBColor(0x3D, 0x9C, 0xF0)
PRED = RGBColor(0xFF, 0x6F, 0x00)
PREY = RGBColor(0x00, 0xBC, 0xD4)
OK = RGBColor(0x66, 0xBB, 0x6A)
WARN = RGBColor(0xFF, 0xC1, 0x07)
BORDER = RGBColor(0x2A, 0x3A, 0x4D)


def set_run(run, size=18, bold=False, color=TEXT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def fill_bg(slide):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def add_title(slide, text, top=0.35, size=32, color=ACCENT):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(top), Inches(12.1), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=True, color=color)


def add_text(slide, text, left, top, width, height, size=18, bold=False, color=TEXT):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for line in text.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)


def add_bullets(slide, items, left, top, width, height, size=18, color=TEXT):
    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)


def card(slide, left, top, width, height):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = PANEL
    shape.line.color.rgb = BORDER
    shape.adjustments[0] = 0.1


def add_img(slide, name, left, top, width=None):
    path = ASSETS / name
    if not path.exists():
        return
    if width:
        slide.shapes.add_picture(
            str(path), Inches(left), Inches(top), width=Inches(width)
        )
    else:
        slide.shapes.add_picture(str(path), Inches(left), Inches(top))


def blank():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill_bg(slide)
    return slide


# 1 Title
s = blank()
add_text(s, "CSE 3812  ·  AI Lab", 0.6, 1.6, 12, 0.4, size=16, color=MUTED)
add_text(s, "Predator & Prey", 0.6, 2.1, 12, 0.9, size=44, bold=True)
add_text(
    s,
    "Grid-based tactical AI simulation — predators hunt, prey forage, "
    "win by elimination or resource control.",
    0.6,
    3.1,
    11,
    0.7,
    size=18,
    color=MUTED,
)
for i, (t, d) in enumerate(
    [
        ("Search", "BFS · A*"),
        ("Behavior", "FSM · Intents · Fcost"),
        ("Core", "Turns · Combat · Occupancy"),
        ("Vision", "Full awareness (v1)"),
    ]
):
    x = 0.6 + i * 3.1
    card(s, x, 4.3, 2.9, 1.5)
    add_text(s, t, x + 0.2, 4.45, 2.5, 0.4, size=16, bold=True, color=ACCENT)
    add_text(s, d, x + 0.2, 4.95, 2.5, 0.6, size=15)

# 2 Problem
s = blank()
add_title(s, "Problem")
card(s, 0.6, 1.4, 5.8, 3.8)
add_text(s, "PREDATORS", 0.9, 1.6, 5, 0.35, size=14, bold=True, color=PRED)
add_text(s, "Eliminate ≥ 70% of prey", 0.9, 2.2, 5.2, 0.6, size=26, bold=True)
add_text(
    s,
    "Tiger (one-shot)  ·  Wolves (damage + cooldown)",
    0.9,
    3.2,
    5.2,
    0.8,
    size=16,
    color=MUTED,
)
card(s, 6.9, 1.4, 5.8, 3.8)
add_text(s, "PREY", 7.2, 1.6, 5, 0.35, size=14, bold=True, color=PREY)
add_text(s, "Consume ≥ 65% of food", 7.2, 2.2, 5.2, 0.6, size=26, bold=True)
add_text(
    s,
    "Deer herds  ·  shared goals  ·  flee under threat",
    7.2,
    3.2,
    5.2,
    0.8,
    size=16,
    color=MUTED,
)
add_text(
    s,
    "Alternating team turns  ·  Predators first each round  ·  20×20 map",
    0.6,
    5.6,
    12,
    0.5,
    size=16,
    color=MUTED,
)

# 3 Architecture
s = blank()
add_title(s, "Architecture")
add_text(
    s,
    "config/*.json  →  data/loader  →  core (rules)  →  ai (decisions)  →  rendering",
    0.6,
    1.3,
    12,
    0.5,
    size=18,
    bold=True,
    color=ACCENT,
)
for i, (t, d) in enumerate(
    [
        (
            "core/",
            "Grid · Agents · Combat · Simulation · Scenario\n"
            "No Pygame — headless + tested",
        ),
        ("ai/", "BFS · A* · Fcost · FSM · FoV (ready) · Influence (stub)"),
        ("rendering/", "Pygame view only — never mutates simulation"),
    ]
):
    x = 0.6 + i * 4.15
    card(s, x, 2.2, 3.95, 3.5)
    add_text(s, t, x + 0.25, 2.45, 3.4, 0.4, size=20, bold=True, color=ACCENT)
    add_text(s, d, x + 0.25, 3.1, 3.4, 2.2, size=16)

# 4 Map
s = blank()
add_title(s, "Map & Units (live render)")
add_img(s, "01_start.png", 0.4, 1.2, width=8.2)
add_bullets(
    s,
    [
        "River + 2 chokepoints",
        "3 feeding grounds · resource nodes",
        "Rocks block movement",
        "T Tiger · W Wolves",
        "D 6 deer · 3 herds",
    ],
    8.9,
    1.5,
    4,
    5,
    size=18,
)

# 5 Workflow
s = blank()
add_title(s, "Core Workflow — One Round")
add_text(
    s,
    "Predator turn  →  move + attack  →  Prey turn  →  move + feed  →  Win check",
    0.6,
    1.25,
    12,
    0.5,
    size=17,
    bold=True,
    color=ACCENT,
)
card(s, 0.6, 2.1, 5.9, 3.8)
add_text(s, "HARD RULES", 0.9, 2.3, 5, 0.35, size=13, bold=True, color=MUTED)
add_bullets(
    s,
    [
        "No shared tiles (occupancy)",
        "Moves sequential · adjacent only",
        "Predators attack from next tile — never stack on prey",
    ],
    0.9,
    2.9,
    5.3,
    2.5,
    size=17,
)
card(s, 6.8, 2.1, 5.9, 3.8)
add_text(s, "DYNAMIC EACH TURN", 7.1, 2.3, 5, 0.35, size=13, bold=True, color=MUTED)
add_bullets(
    s,
    [
        "Positions · HP · food stock",
        "Prey state · wolf cooldown",
        "Flock food goal · events · winner",
    ],
    7.1,
    2.9,
    5.3,
    2.5,
    size=17,
)

# 6 Midgame
s = blank()
add_title(s, "Mid-Game State")
add_img(s, "02_midgame.png", 0.4, 1.2, width=8.2)
add_bullets(
    s,
    [
        "Herds split across feeding areas",
        "Brown tiles = depleted food",
        "Same rules in Fast & Detailed playback",
        "Screenshot from real Renderer + Simulation",
    ],
    8.9,
    1.8,
    4,
    4.5,
    size=17,
)

# 7 Algorithms
s = blank()
add_title(s, "Algorithms — Search")
card(s, 0.6, 1.3, 12.1, 3.6)
add_text(s, "Algo", 0.9, 1.5, 2.5, 0.4, size=14, bold=True, color=MUTED)
add_text(s, "Where", 3.5, 1.5, 4, 0.4, size=14, bold=True, color=MUTED)
add_text(s, "Does", 8.0, 1.5, 4, 0.4, size=14, bold=True, color=MUTED)
for i, (a, w, d) in enumerate(
    [
        ("BFS", "Nearest food · flock goal", "Unweighted shortest reach"),
        ("A*", "Prey NORMAL · predator intents", "Path with Chebyshev heuristic"),
        ("Chebyshev", "Everywhere", "8-dir distance · adjacency"),
    ]
):
    y = 2.1 + i * 0.85
    add_text(s, a, 0.9, y, 2.5, 0.5, size=18, bold=True, color=ACCENT)
    add_text(s, w, 3.5, y, 4.2, 0.5, size=17)
    add_text(s, d, 8.0, y, 4.3, 0.5, size=17)
card(s, 0.6, 5.2, 12.1, 1.3)
add_text(
    s,
    "A*:   f = g + h     |     h = max(|Δr|, |Δc|)",
    0.9,
    5.35,
    11,
    0.4,
    size=20,
    bold=True,
)
add_text(
    s,
    "Take only the next step each turn — replan next round.",
    0.9,
    5.9,
    11,
    0.4,
    size=15,
    color=MUTED,
)

# 8 Intents
s = blank()
add_title(s, "Predator Intents (rolled once)")
for i, (t, d, note) in enumerate(
    [
        ("Chase", "A* toward closest prey", "Also after tiger's first kill"),
        ("Ambush", "A* to midpoint(prey, prey's food)", "Wolf bias"),
        (
            "Camp",
            "Wait opposite side of river",
            "If prey < 3 tiles → close in · Tiger bias",
        ),
    ]
):
    x = 0.6 + i * 4.15
    card(s, x, 1.5, 3.95, 3.6)
    add_text(s, t, x + 0.25, 1.75, 3.4, 0.5, size=22, bold=True, color=ACCENT)
    add_text(s, d, x + 0.25, 2.5, 3.4, 1.2, size=17)
    add_text(s, note, x + 0.25, 4.0, 3.4, 0.8, size=15, color=MUTED)
add_text(
    s,
    "Target = closest living prey · recomputed every move",
    0.6,
    5.5,
    12,
    0.5,
    size=18,
)

# 9 Prey FSM
s = blank()
add_title(s, "Prey FSM")
add_text(
    s,
    "NORMAL  →  PANIC  →  DESPAIR  →  NORMAL if predator > 3",
    0.6,
    1.25,
    12,
    0.45,
    size=18,
    bold=True,
    color=ACCENT,
)
for i, (t, d) in enumerate(
    [
        ("NORMAL", "Shared flock food goal · A* · eat underfoot if safe"),
        ("PANIC", "Flee + still prefer food (fcost)"),
        ("DESPAIR", "Flee + regroup flock · no food term"),
    ]
):
    x = 0.6 + i * 4.15
    card(s, x, 2.0, 3.95, 3.0)
    add_text(s, t, x + 0.25, 2.25, 3.4, 0.45, size=20, bold=True, color=PREY)
    add_text(s, d, x + 0.25, 2.95, 3.4, 1.6, size=16)
add_text(
    s,
    "Despair triggers on kill · Panic scoring implemented · spotting trigger = Phase 2",
    0.6,
    5.4,
    12,
    0.5,
    size=15,
    color=MUTED,
)

# 10 Fcost
s = blank()
add_title(s, "Fcost ≠ A*'s f")
card(s, 0.6, 1.4, 5.9, 3.6)
add_text(s, "WHAT IT IS", 0.9, 1.6, 5, 0.35, size=13, bold=True, color=MUTED)
add_text(s, "Local rank score over ≤8 neighbors (+ stay)", 0.9, 2.2, 5.3, 0.7, size=18)
add_text(
    s,
    "best gets N … worst gets 1\nfcost = sum of ranks",
    0.9,
    3.2,
    5.3,
    1.2,
    size=18,
    bold=True,
    color=ACCENT,
)
card(s, 6.8, 1.4, 5.9, 3.6)
add_text(s, "FORMULAS", 7.1, 1.6, 5, 0.35, size=13, bold=True, color=MUTED)
add_text(s, "Panic: away from predator + toward food", 7.1, 2.3, 5.3, 0.8, size=17)
add_text(
    s,
    "Despair: away from predator + toward mate (unless already ≤2)",
    7.1,
    3.3,
    5.3,
    1.0,
    size=17,
)
add_text(
    s,
    "A* = global path to a goal    ·    Fcost = one emergency step under fear",
    0.6,
    5.4,
    12,
    0.5,
    size=17,
    color=MUTED,
)

# 11 Combat
s = blank()
add_title(s, "Combat & Events")
add_img(s, "04_kill_event.png", 0.4, 1.2, width=8.0)
add_bullets(
    s,
    [
        "Adjacent after move → attack",
        "Tiger one-shots · Wolf 60 dmg",
        "Wolf skips attacks 1 full turn after kill",
        "KILL / EMPTY FX + sidebar log",
    ],
    8.7,
    1.8,
    4.2,
    4.5,
    size=17,
)

# 12 Vision
s = blank()
add_title(s, "Vision")
card(s, 0.6, 1.3, 12.1, 1.6)
add_text(
    s,
    "v1 = Full awareness — whole map known",
    0.9,
    1.55,
    11.5,
    0.5,
    size=24,
    bold=True,
)
add_text(
    s,
    "No fog of war in the live turn loop",
    0.9,
    2.2,
    11.5,
    0.4,
    size=16,
    color=MUTED,
)
card(s, 0.6, 3.3, 5.9, 2.8)
add_text(s, "READY, NOT WIRED", 0.9, 3.5, 5, 0.35, size=13, bold=True, color=WARN)
add_bullets(
    s,
    ["src/ai/fov.py", "Rocks block Bresenham LoS", "Tall grass hides past 3 tiles"],
    0.9,
    4.1,
    5.3,
    1.8,
    size=16,
)
card(s, 6.8, 3.3, 5.9, 2.8)
add_text(s, "DEFEND", 7.1, 3.5, 5, 0.35, size=13, bold=True, color=OK)
add_text(
    s,
    "Deterministic chase & tests first. LoS is an optional harder mode per GDD.",
    7.1,
    4.2,
    5.3,
    1.5,
    size=17,
)

# 13 CSP
s = blank()
add_title(s, "CSP?")
card(s, 0.6, 1.3, 12.1, 1.4)
add_text(
    s,
    "Not used — current AI = search + FSM + local scoring",
    0.9,
    1.65,
    11.5,
    0.6,
    size=22,
    bold=True,
    color=WARN,
)
card(s, 0.6, 3.1, 5.9, 3.0)
add_text(s, "WHERE CSP COULD FIT", 0.9, 3.3, 5, 0.35, size=13, bold=True, color=MUTED)
add_bullets(
    s,
    [
        "Legal spawn layouts",
        "Simultaneous herd next-tiles",
        "Wolf pack spacing assignment",
    ],
    0.9,
    3.9,
    5.3,
    2,
    size=17,
)
card(s, 6.8, 3.1, 5.9, 3.0)
add_text(s, "WHAT WE DO INSTEAD", 7.1, 3.3, 5, 0.35, size=13, bold=True, color=OK)
add_text(
    s,
    "Sequential moves + hard occupancy checks = constraints enforced, not a CSP solver.",
    7.1,
    4.1,
    5.3,
    1.6,
    size=17,
)

# 14 Detailed
s = blank()
add_title(s, "Detailed Playback")
add_img(s, "03_detailed_action.png", 0.4, 1.2, width=8.0)
add_bullets(
    s,
    [
        "One unit action per tick",
        "Active unit highlighted",
        "Fast mode = full round / tick",
        "Identical outcomes (tested)",
    ],
    8.7,
    1.8,
    4.2,
    4.5,
    size=17,
)

# 15 Feeding
s = blank()
add_title(s, "Foraging & Depletion")
add_img(s, "05_feeding_depleted.png", 0.4, 1.2, width=8.0)
add_bullets(
    s,
    [
        "Eat underfoot when safe",
        "Blocked goal → BFS other food",
        "Depleted = brown · still passable",
        "Prey win at 65% pool eaten",
    ],
    8.7,
    1.8,
    4.2,
    4.5,
    size=17,
)

# 16 Debug
s = blank()
add_title(s, "Debug Overlay")
add_img(s, "06_debug_overlay.png", 0.4, 1.2, width=8.0)
add_bullets(
    s,
    [
        "F1 — intents + prey states",
        "Shows Chase / Ambush / Camp",
        "Shows N / P / D state letter",
    ],
    8.7,
    2.0,
    4.2,
    4,
    size=17,
)

# 17 End
s = blank()
add_title(s, "End State")
add_img(s, "07_endgame.png", 0.4, 1.2, width=8.0)
add_bullets(
    s,
    [
        "Winner set when threshold hit",
        "Sample balance ~55% prey / 45% pred (seeds 0–99)",
        "History logged each round",
    ],
    8.7,
    2.0,
    4.2,
    4,
    size=17,
)

# 18 Gaps
s = blank()
add_title(s, "Gaps & Defense")
for i, (t, d, note) in enumerate(
    [
        (
            "Panic spotting",
            "Fcost ready · trigger not wired",
            "Phase 2 — despair already on kill",
        ),
        (
            "Wolf pack AI",
            "Coded · default uses solo A*",
            "Next: assign pack_id + call pack step",
        ),
        ("LoS / influence / CSP", "Optional / deferred", "Not required for core demo"),
    ]
):
    x = 0.6 + i * 4.15
    card(s, x, 1.4, 3.95, 3.2)
    add_text(s, t, x + 0.2, 1.6, 3.5, 0.45, size=18, bold=True, color=WARN)
    add_text(s, d, x + 0.2, 2.3, 3.5, 0.9, size=15, color=MUTED)
    add_text(s, note, x + 0.2, 3.4, 3.5, 0.9, size=15)
card(s, 0.6, 5.0, 12.1, 1.5)
add_text(
    s,
    "DONE: Core loop · occupancy · BFS/A* · intents · despair · fcost · "
    "combat · wins · playback · tests",
    0.9,
    5.4,
    11.5,
    0.7,
    size=17,
    bold=True,
    color=OK,
)

# 19 Controls
s = blank()
add_title(s, "Demo Controls")
card(s, 2.5, 1.3, 8.3, 4.5)
for i, (k, a) in enumerate(
    [
        ("Space", "Pause / resume"),
        ("T", "Detailed ↔ Fast"),
        ("N", "Next action / round"),
        ("+ / −", "Speed"),
        ("R", "New scenario"),
        ("F1", "Debug overlay"),
    ]
):
    y = 1.55 + i * 0.6
    add_text(s, k, 3.0, y, 2.5, 0.45, size=18, bold=True, color=ACCENT)
    add_text(s, a, 5.8, y, 4.5, 0.45, size=18)
add_text(
    s,
    "python -m src.main    ·    pytest tests/ -v",
    0.6,
    6.2,
    12,
    0.4,
    size=16,
    color=MUTED,
)

# 20 Summary
s = blank()
add_title(s, "Summary", color=TEXT)
add_bullets(
    s,
    [
        "Headless core owns rules; AI chooses moves; render only shows",
        "Search: BFS + A* · Fear: ranked fcost · Behavior: FSM + intents",
        "Vision: full map in v1 · FoV module ready for later",
        "Honest Phase-1 complete · Phase-2 wires identified",
    ],
    0.8,
    1.5,
    11.5,
    3.5,
    size=22,
)
add_text(s, "Questions?", 0.8, 5.5, 11, 0.6, size=28, bold=True, color=ACCENT)

prs.save(str(OUT))
print(f"Wrote {OUT}")
print(f"Slides: {len(prs.slides)}")
